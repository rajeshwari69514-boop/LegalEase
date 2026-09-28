import streamlit as st
import requests
from docx import Document
from io import BytesIO
from fpdf import FPDF

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="centered"
)

st.markdown(
    """
    <style>
    .main-title {
        text-align: center;
        font-size: 42px;
        font-weight: 700;
        margin-bottom: 5px;
    }

    .subtitle {
        text-align: center;
        font-size: 18px;
        margin-bottom: 25px;
    }

    .section-title {
        font-size: 24px;
        font-weight: 600;
        margin-top: 20px;
        margin-bottom: 15px;
    }

    .stButton > button {
        width: 100%;
        height: 50px;
        font-size: 18px;
        font-weight: 600;
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.markdown(
    '<div class="main-title">⚖️ LegalEase</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">AI-Powered Legal Document Generator</div>',
    unsafe_allow_html=True
)

st.info(
    "Create professional legal document drafts quickly and easily."
)

st.divider()

st.markdown(
    '<div class="section-title">📝 Document Details</div>',
    unsafe_allow_html=True
)

document_type = st.selectbox(
    "Select Document Type",
    [
        "NDA",
        "Employment Contract",
        "Lease Agreement",
        "Service Agreement"
    ]
)

purpose = st.text_area(
    "Purpose",
    placeholder="Enter the purpose of this document"
)

date = st.text_input(
    "Date",
    placeholder="DD-MM-YYYY"
)

st.markdown(
    '<div class="section-title">👤 Party Details</div>',
    unsafe_allow_html=True
)

col1, col2 = st.columns(2)

with col1:
    st.markdown("### Party 1")
    party1_name = st.text_input(
        "Party 1 Name",
        placeholder="Enter name",
        key="party1"
    )

with col2:
    st.markdown("### Party 2")
    party2_name = st.text_input(
        "Party 2 Name",
        placeholder="Enter name",
        key="party2"
    )

st.divider()

generate = st.button(
    "✨ Generate Legal Document",
    type="primary"
)

if generate:

    if not party1_name or not party2_name or not purpose or not date:
        st.warning(
            "Please fill in all the required details."
        )

    else:

        data = {
            "document_type": document_type,
            "party1_name": party1_name,
            "party2_name": party2_name,
            "purpose": purpose,
            "date": date
        }

        try:

            with st.spinner("Generating your legal document..."):

                response = requests.post(
                    "http://127.0.0.1:8000/generate",
                    json=data
                )

            if response.status_code == 200:

                result = response.json()
                document_text = result["document"]

                st.divider()

                st.markdown(
                    '<div class="section-title">📄 Generated Document</div>',
                    unsafe_allow_html=True
                )

                st.text_area(
                    "Document Preview",
                    document_text,
                    height=500
                )

                # =========================
                # DOCX
                # =========================

                doc = Document()

                title = doc.add_heading(
                    document_type,
                    level=1
                )

                title.alignment = 1

                for paragraph in document_text.split("\n"):

                    paragraph = paragraph.strip()

                    if paragraph:

                        p = doc.add_paragraph()

                        p.paragraph_format.space_after = 8
                        p.paragraph_format.line_spacing = 1.15

                        p.add_run(paragraph)

                docx_file = BytesIO()

                doc.save(docx_file)

                docx_file.seek(0)

                # =========================
                # PDF SAFE TEXT
                # =========================

                pdf_text = (
                    document_text
                    .replace("“", '"')
                    .replace("”", '"')
                    .replace("‘", "'")
                    .replace("’", "'")
                    .replace("–", "-")
                    .replace("—", "-")
                    .replace("•", "-")
                    .replace("…", "...")
                )

                pdf_text = (
                    pdf_text
                    .encode("ascii", "ignore")
                    .decode("ascii")
                )

                # =========================
                # PDF
                # =========================

                pdf = FPDF()

                pdf.set_auto_page_break(
                    auto=True,
                    margin=20
                )

                pdf.set_margins(
                    20,
                    20,
                    20
                )

                pdf.add_page()

                pdf.set_font(
                    "Helvetica",
                    "B",
                    18
                )

                pdf.cell(
                    0,
                    12,
                    document_type,
                    align="C"
                )

                pdf.ln(15)

                for paragraph in pdf_text.split("\n"):

                    paragraph = paragraph.strip()

                    if not paragraph:

                        pdf.ln(4)

                        continue

                    is_heading = (
                        paragraph.isupper()
                        or paragraph.startswith("1.")
                        or paragraph.startswith("2.")
                        or paragraph.startswith("3.")
                        or paragraph.startswith("4.")
                        or paragraph.startswith("5.")
                        or paragraph.startswith("6.")
                        or paragraph.startswith("7.")
                        or paragraph.startswith("8.")
                        or paragraph.startswith("9.")
                        or paragraph.startswith("10.")
                    )

                    if is_heading:

                        pdf.set_font(
                            "Helvetica",
                            "B",
                            12
                        )

                        pdf.multi_cell(
                            0,
                            8,
                            paragraph
                        )

                        pdf.ln(2)

                    else:

                        pdf.set_font(
                            "Helvetica",
                            "",
                            11
                        )

                        pdf.multi_cell(
                            0,
                            7,
                            paragraph
                        )

                        pdf.ln(3)

                pdf_file = bytes(
                    pdf.output()
                )

                # =========================
                # DOWNLOAD BUTTONS
                # =========================

                st.markdown(
                    "### 📥 Download Document"
                )

                download_col1, download_col2 = st.columns(2)

                with download_col1:

                    st.download_button(
                        label="📄 Download DOCX",
                        data=docx_file,
                        file_name=f"{document_type}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )

                with download_col2:

                    st.download_button(
                        label="📑 Download PDF",
                        data=pdf_file,
                        file_name=f"{document_type}.pdf",
                        mime="application/pdf",
                        use_container_width=True
                    )

                st.success(
                    "Your legal document has been generated successfully!"
                )

                st.caption(
                    "Disclaimer: This document is an AI-generated general draft and is not a substitute for professional legal advice."
                )

            else:

                st.error(
                    f"Something went wrong. Status code: {response.status_code}"
                )

                st.write(response.text)

        except Exception as e:

            st.error("Connection error")

            st.write(str(e))